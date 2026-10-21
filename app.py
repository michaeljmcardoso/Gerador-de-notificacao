from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PROCESS_NUMBER = "54230.006073/2006-24"
NOTIFICATION_NUMBER = "00000000/2026/SR(12)MA-Q1/SR(12)MA-Q/SR(12)MA/INCRA-INCRA"
TERRITORY = "Bom Jesus/Santa Maria dos Moreiras/Jerusalém"
INCIDENCE_OPTIONS = (
    "Incide parcialmente",
    "Incide totalmente",
    "É confrontante",
)
INCIDENCE_TEXT = {
    "Incide parcialmente": "incide parcialmente no perímetro",
    "Incide totalmente": "incide totalmente no perímetro",
    "É confrontante": "é confrontante ao",
}
LOGO_PATH = Path(__file__).parent / "assets" / "brasao.jpg"
DOCUMENT_LANGUAGE = "pt-BR"


@dataclass
class NotificationData:
    representative: str
    property_name: str
    registry: str
    header_code: str
    address: str
    body_code: str
    incidence: str


def _set_run_language(run):
    properties = run._element.get_or_add_rPr()
    language = properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        properties.append(language)
    language.set(qn("w:val"), DOCUMENT_LANGUAGE)


def _set_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    _set_run_language(run)
    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_page_number(paragraph):
    run = paragraph.add_run()
    _set_font(run, size=8, color=(160, 160, 160))
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def _add_body_paragraph(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    _set_font(run)
    return paragraph


def build_notification_docx(data: NotificationData) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)
    section.footer_distance = Cm(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal_properties = normal._element.get_or_add_rPr()
    normal_language = normal_properties.find(qn("w:lang"))
    if normal_language is None:
        normal_language = OxmlElement("w:lang")
        normal_properties.append(normal_language)
    normal_language.set(qn("w:val"), DOCUMENT_LANGUAGE)

    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(2)
    header.add_run().add_picture(str(LOGO_PATH), width=Cm(1.8))

    for text in (
        "INSTITUTO NACIONAL DE COLONIZAÇÃO E REFORMA AGRÁRIA",
        "Rua H, Quadra E, nº. 12 - Jardim Atlântico, INCRA - Bairro Turu, São Luís/MA, CEP 65067-150",
        "(98) 3878-7450",
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        _set_font(run, size=9 if text.startswith("Rua ") else 10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(25)
    run = title.add_run(f"MINUTA DE NOTIFICAÇÃO Nº {NOTIFICATION_NUMBER}")
    _set_font(run, size=11)

    recipient_lines = [
        "Aos Representantes/Titulares",
        data.representative,
        data.property_name,
        f"Matrícula {data.registry}",
    ]
    if data.header_code:
        recipient_lines.append(f"Código do Imóvel {data.header_code}")
    recipient_lines.append(data.address)
    for text in recipient_lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        _set_font(run)

    subject = document.add_paragraph()
    subject.paragraph_format.space_before = Pt(18)
    subject.paragraph_format.space_after = Pt(0)
    subject.add_run("Assunto: ")
    run = subject.add_run(
        "Notificação sobre publicação do Relatório Técnico de Identificação e "
        f"Delimitação - RTID - do território quilombola {TERRITORY}"
    )
    _set_font(run, bold=True)

    reference = document.add_paragraph()
    reference.paragraph_format.space_after = Pt(16)
    reference.add_run(
        "Referência: Caso responda este Ofício, indicar expressamente o "
        f"Processo nº {PROCESS_NUMBER}."
    )

    greeting = document.add_paragraph()
    greeting.alignment = WD_ALIGN_PARAGRAPH.CENTER
    greeting.paragraph_format.space_after = Pt(5)
    greeting.add_run("Prezado Senhor,")

    _add_body_paragraph(
        document,
        "O Instituto Nacional de Colonização e Reforma Agrária – INCRA, por meio "
        "da Superintendência Regional do Estado do Maranhão, comunica, com "
        "fundamento no art. 68 do Ato das Disposições Constitucionais Transitórias "
        "da Constituição Federal de 1988, no Decreto n.º 4.887/2003 e na Instrução "
        "Normativa INCRA n.º 57/2009, que tramita nesta Superintendência o Processo "
        f"Administrativo n.º {PROCESS_NUMBER}, que trata da regularização fundiária "
        "das terras das Comunidades Remanescentes de Quilombo "
        f"{TERRITORY}, localizadas no Município de Codó, Estado do Maranhão.",
    )

    relation = INCIDENCE_TEXT[data.incidence]
    incidence_clause = (
        f"{relation} território quilombola de {TERRITORY}"
        if data.incidence == "É confrontante"
        else f"{relation} do território quilombola de {TERRITORY}"
    )
    body_code_text = f", código {data.body_code}" if data.body_code else ""
    _add_body_paragraph(
        document,
        "Em cumprimento ao estabelecido no citado Decreto e na Instrução Normativa, "
        "foi elaborado o Relatório Técnico de Identificação e Delimitação – RTID – "
        "do território das referidas Comunidades, cujo extrato foi publicado por "
        "meio de edital no Diário Oficial da União (conforme cópia anexa). Foi "
        f"constatado que o imóvel rural denominado {data.property_name}"
        f"{body_code_text}, matrícula n.º {data.registry}, que se encontra sob sua "
        f"posse/propriedade, {incidence_clause}.",
    )

    _add_body_paragraph(
        document,
        "Assim, em conformidade com o § 2º do art. 7º, combinado com o art. 9º do "
        "aludido Decreto, e com o § 2º do art. 11 da Instrução Normativa citada, "
        "fica o senhor notificado para, querendo, apresentar contestação ao RTID "
        "no prazo de 90 (noventa) dias, contados a partir do recebimento desta "
        "notificação, juntando as provas pertinentes.",
    )

    _add_body_paragraph(
        document,
        "Informa ainda que o Processo Administrativo n.º "
        f"{PROCESS_NUMBER}, em cujos autos se processa o feito, estará à sua "
        "disposição para consulta na sede desta Superintendência, situada na Rua H, "
        "Quadra E, Lote 01, n.º 12 – Bairro Turu – São Luís/MA, de segunda a "
        "sexta-feira, durante o expediente das 8h às 12h e das 14h às 18h, ou pelos "
        "canais de comunicação indicados no rodapé deste documento.",
    )

    closing = document.add_paragraph()
    closing.paragraph_format.space_before = Pt(5)
    closing.paragraph_format.space_after = Pt(0)
    run = closing.add_run("Atenciosamente,")
    _set_font(run)

    signature = document.add_paragraph()
    signature.paragraph_format.space_before = Pt(16)
    signature.paragraph_format.space_after = Pt(0)
    _set_font(signature.add_run("____________________________"))

    for text in (
        "E-mail: divisao.quilombolas.ma@incra.gov.br",
        "Telefone: 98 3210-7474 | 7479",
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        _set_font(paragraph.add_run(text), size=9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Minuta de Notificação 00000000     SEI ")
    footer.add_run(f"{PROCESS_NUMBER} / pg. ")
    _add_page_number(footer)

    for paragraph in (
        *document.paragraphs,
        *section.header.paragraphs,
        *section.footer.paragraphs,
    ):
        for run in paragraph.runs:
            _set_run_language(run)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


st.set_page_config(page_title="Gerador de Minuta de Notificação", page_icon="📄")
st.title("Gerador de Minuta de Notificação")
st.write(
    "Preencha os dados do representante e do imóvel. Os textos institucionais "
    "seguem o modelo de notificação padrão."
)

with st.form("notification_form"):
    st.subheader("Destinatário e imóvel")
    representative = st.text_input("Nome do representante/titular *")
    property_name = st.text_input("Nome do imóvel *")
    registry = st.text_input("Matrícula do imóvel *")
    header_code = st.text_input("Código do imóvel no cabeçalho (opcional)")
    address = st.text_area("Endereço *", height=90)

    st.subheader("Informações do território")
    st.caption(
        "O código no parágrafo é independente do código do cabeçalho, como no "
        "modelo fornecido."
    )
    body_code = st.text_input("Código do imóvel no segundo parágrafo (opcional)")
    incidence = st.selectbox("Relação com o território quilombola", INCIDENCE_OPTIONS)
    submitted = st.form_submit_button("Gerar minuta DOCX")

if submitted:
    required_fields = {
        "Nome do representante/titular": representative,
        "Nome do imóvel": property_name,
        "Matrícula do imóvel": registry,
        "Endereço": address,
    }
    missing_fields = [label for label, value in required_fields.items() if not value.strip()]
    if missing_fields:
        st.error("Preencha os campos obrigatórios: " + ", ".join(missing_fields) + ".")
    else:
        data = NotificationData(
            representative=representative.strip(),
            property_name=property_name.strip(),
            registry=registry.strip(),
            header_code=header_code.strip(),
            address=address.strip(),
            body_code=body_code.strip(),
            incidence=incidence,
        )
        document_bytes = build_notification_docx(data)
        st.success("Minuta gerada.")
        st.download_button(
            "Baixar minuta em DOCX",
            data=document_bytes,
            file_name="minuta_de_notificacao.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
