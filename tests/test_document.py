import unittest
from collections import Counter
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from app import (
    DATABASE_PATH,
    NotificationData,
    ProcessRecord,
    build_notification_docx,
    format_process_record,
    load_process_records,
)


class NotificationDocumentTests(unittest.TestCase):
    def setUp(self):
        self.data = NotificationData(
            representative="Representante Exemplo",
            property_name="Fazenda Exemplo",
            registry="12345",
            header_code="1111111111111",
            address="Estrada Exemplo, Codó-MA",
            body_code="2222222222222",
            incidence="Incide parcialmente",
            process_number="54230.006073/2006-24",
            community="Bom Jesus/Santa Maria dos Moreiras/Jerusalém",
            municipality="Codó",
        )

    def get_text(self, incidence=None):
        data = self.data
        if incidence:
            data = NotificationData(
                representative=data.representative,
                property_name=data.property_name,
                registry=data.registry,
                header_code=data.header_code,
                address=data.address,
                body_code=data.body_code,
                incidence=incidence,
                process_number=data.process_number,
                community=data.community,
                municipality=data.municipality,
            )
        document = Document(BytesIO(build_notification_docx(data)))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def test_generates_docx_with_distinct_header_and_body_codes(self):
        text = self.get_text()
        self.assertIn("Código do Imóvel SNCR 1111111111111", text)
        self.assertIn("código 2222222222222, matrícula n.º 12345", text)
        self.assertIn("incide parcialmente no perímetro", text)
        self.assertIn("Processo nº 54230.006073/2006-24", text)

    def test_codes_are_optional_and_omitted_when_empty(self):
        data = NotificationData(
            representative=self.data.representative,
            property_name=self.data.property_name,
            registry=self.data.registry,
            header_code="",
            address=self.data.address,
            body_code="",
            incidence=self.data.incidence,
            process_number=self.data.process_number,
            community=self.data.community,
            municipality=self.data.municipality,
        )
        document = Document(BytesIO(build_notification_docx(data)))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertNotIn("Código do Imóvel", text)
        self.assertNotIn(", código ", text)
        self.assertIn(
            "imóvel rural denominado Fazenda Exemplo, matrícula n.º 12345",
            text,
        )

    def test_community_process_and_municipality_are_used_in_document(self):
        text = self.get_text()
        self.assertIn(
            "Processo nº 54230.006073/2006-24",
            text,
        )
        self.assertIn(
            "Comunidade Remanescente de Quilombo "
            "Bom Jesus/Santa Maria dos Moreiras/Jerusalém, localizada no Município "
            "de Codó",
            text,
        )
        self.assertIn(
            "território quilombola de Bom Jesus/Santa Maria dos Moreiras/Jerusalém",
            text,
        )

    def test_codes_can_be_filled_independently(self):
        data = NotificationData(
            representative=self.data.representative,
            property_name=self.data.property_name,
            registry=self.data.registry,
            header_code="",
            address=self.data.address,
            body_code=self.data.body_code,
            incidence=self.data.incidence,
            process_number=self.data.process_number,
            community=self.data.community,
            municipality=self.data.municipality,
        )
        text = "\n".join(
            paragraph.text
            for paragraph in Document(BytesIO(build_notification_docx(data))).paragraphs
        )
        self.assertNotIn("Código do Imóvel", text)
        self.assertIn("código 2222222222222, matrícula n.º 12345", text)

    def test_incidence_options_use_their_matching_phrases(self):
        expected = {
            "Incide parcialmente": "incide parcialmente no perímetro",
            "Incide totalmente": "incide totalmente no perímetro",
            "É confrontante": "é confrontante ao território",
        }
        for incidence, phrase in expected.items():
            with self.subTest(incidence=incidence):
                self.assertIn(phrase, self.get_text(incidence))

    def test_standard_paragraphs_and_footer_are_present(self):
        document = Document(BytesIO(build_notification_docx(self.data)))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        footer = document.sections[0].footer.paragraphs[0].text

        self.assertIn("art. 68 do Ato das Disposições Constitucionais Transitórias", text)
        self.assertIn("prazo de 90 (noventa) dias", text)
        self.assertIn("de segunda a sexta-feira", text)
        self.assertIn("divisao.quilombolas.ma@incra.gov.br", text)
        self.assertIn("SEI 54230.006073/2006-24", footer)

    def test_process_database_preserves_all_complete_records_including_duplicates(self):
        records = load_process_records(str(DATABASE_PATH))
        self.assertEqual(len(records), 445)
        self.assertEqual(len(records), len(set(records)))
        self.assertEqual(
            [record.record_id for record in records],
            sorted(record.record_id for record in records),
        )
        self.assertTrue(
            all(record.process_number and record.community and record.municipality for record in records)
        )
        data_counts = Counter(
            (record.process_number, record.community, record.municipality)
            for record in records
        )
        self.assertEqual(sum(count > 1 for count in data_counts.values()), 2)
        self.assertEqual(sum(data_counts.values()), 445)
        self.assertEqual(
            data_counts[("54230.000952/2017-03", "Santa Rosa", "Pinheiro")],
            2,
        )
        self.assertEqual(
            data_counts[("54230.002398/2006-38", "Mucambo", "Viana")],
            2,
        )

    def test_duplicate_community_names_are_disambiguated_in_selector(self):
        records = (
            ProcessRecord(1, "54230.000001/2020-11", "Comunidade Exemplo", "Município A"),
            ProcessRecord(2, "54230.000002/2020-22", "Comunidade Exemplo", "Município B"),
            ProcessRecord(3, "54230.000003/2020-33", "Comunidade Repetida", "Município C"),
            ProcessRecord(4, "54230.000004/2020-44", "Comunidade Repetida", "Município C"),
        )
        community_counts = Counter(record.community for record in records)
        pair_counts = Counter(
            (record.community, record.municipality) for record in records
        )
        labels = [
            format_process_record(record, community_counts, pair_counts, Counter())
            for record in records
        ]
        label_counts = Counter(labels)

        self.assertEqual(
            format_process_record(records[0], community_counts, pair_counts, label_counts),
            "Comunidade Exemplo — Município A",
        )
        self.assertEqual(
            format_process_record(records[2], community_counts, pair_counts, label_counts),
            "Comunidade Repetida — Município C — Processo 54230.000003/2020-33",
        )

    def test_identical_process_rows_get_unique_selector_labels(self):
        records = (
            ProcessRecord(10, "54000.000001/2025-05", "Achuí II", "Penalva"),
            ProcessRecord(11, "54000.000001/2025-05", "Achuí II", "Penalva"),
        )
        community_counts = Counter(record.community for record in records)
        pair_counts = Counter(
            (record.community, record.municipality) for record in records
        )
        labels = [
            format_process_record(record, community_counts, pair_counts, Counter())
            for record in records
        ]
        label_counts = Counter(labels)

        self.assertEqual(
            [
                format_process_record(record, community_counts, pair_counts, label_counts)
                for record in records
            ],
            ["Achuí II — Penalva — Processo 54000.000001/2025-05 — Registro 10",
             "Achuí II — Penalva — Processo 54000.000001/2025-05 — Registro 11"],
        )

    def test_document_language_is_brazilian_portuguese(self):
        document = Document(BytesIO(build_notification_docx(self.data)))
        self.assertEqual(
            document.styles["Normal"]._element.rPr.find(qn("w:lang")).get(qn("w:val")),
            "pt-BR",
        )
        paragraphs = (
            *document.paragraphs,
            *document.sections[0].header.paragraphs,
            *document.sections[0].footer.paragraphs,
        )
        for paragraph in paragraphs:
            for run in paragraph.runs:
                with self.subTest(text=run.text):
                    language = run._element.rPr.find(qn("w:lang"))
                    self.assertIsNotNone(language)
                    self.assertEqual(language.get(qn("w:val")), "pt-BR")


if __name__ == "__main__":
    unittest.main()
